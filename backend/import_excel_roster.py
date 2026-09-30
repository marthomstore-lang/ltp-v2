import openpyxl
import mysql.connector
import datetime
import re

def format_rut(run_num, dv):
    if not run_num:
        return ""
    run_str = str(run_num).strip()
    dv_str = str(dv).strip().upper() if dv is not None else ""
    
    # Formatear con puntos: 27560302 -> 27.560.302
    if len(run_str) >= 7:
        formatted_num = f"{run_str[:-6]}.{run_str[-6:-3]}.{run_str[-3:]}"
    else:
        formatted_num = run_str
        
    return f"{formatted_num}-{dv_str}" if dv_str else formatted_num

def parse_date(val):
    if not val:
        return None
    if isinstance(val, datetime.datetime) or isinstance(val, datetime.date):
        # Si es 01/01/1900, retornar None
        if val.year == 1900:
            return None
        return val.strftime('%Y-%m-%d')
    if isinstance(val, str):
        val_str = val.strip()
        if '1900' in val_str or not val_str:
            return None
        return val_str
    return None

def import_roster():
    print("🚀 Iniciando importación de 418 estudiantes desde nomina_excel (16).xlsx a MySQL XAMPP...")
    
    # 1. Conectar a MySQL XAMPP
    db = mysql.connector.connect(
        host="127.0.0.1",
        port=3306,
        user="root",
        password="",
        database="ltp_local_db"
    )
    cursor = db.cursor()
    
    # 2. Agregar columnas adicionales a la tabla students si no existen
    alter_queries = [
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS academic_year INT DEFAULT 2026",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS rbd INT DEFAULT 3941",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS cod_tipo_ensenanza INT",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS cod_grado INT",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS letra_curso VARCHAR(10)",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS run_num VARCHAR(20)",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS dv_run VARCHAR(5)",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS gender VARCHAR(10)",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS first_name VARCHAR(150)",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS paternal_surname VARCHAR(150)",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS maternal_surname VARCHAR(150)",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS address TEXT",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS commune VARCHAR(100)",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS commune_code VARCHAR(20)",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS email VARCHAR(255)",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS phone VARCHAR(50)",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS mobile_phone VARCHAR(50)",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS birth_date DATE",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS ethnicity_code INT",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS enrollment_date DATE",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS attendance_percentage DECIMAL(5,2)",
        "ALTER TABLE students ADD COLUMN IF NOT EXISTS final_average DECIMAL(3,1)"
    ]
    
    for q in alter_queries:
        try:
            cursor.execute(q)
        except Exception as e:
            pass
    db.commit()

    # 3. Cargar libro de Excel
    excel_path = r'C:\Users\david\Downloads\nomina_excel (16).xlsx'
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb['nomina_excel (16)']

    inserted_count = 0
    updated_count = 0

    for i in range(2, sheet.max_row + 1):
        year = sheet.cell(row=i, column=1).value
        rbd = sheet.cell(row=i, column=2).value
        cod_tipo_ens = sheet.cell(row=i, column=3).value
        cod_grado = sheet.cell(row=i, column=4).value
        desc_grado_raw = sheet.cell(row=i, column=5).value
        letra_curso = str(sheet.cell(row=i, column=6).value or 'A').strip()
        run_num = sheet.cell(row=i, column=7).value
        dv_run = sheet.cell(row=i, column=8).value
        gender = sheet.cell(row=i, column=9).value
        nombres = str(sheet.cell(row=i, column=10).value or '').strip()
        ap_paterno = str(sheet.cell(row=i, column=11).value or '').strip()
        ap_materno = str(sheet.cell(row=i, column=12).value or '').strip()
        direccion = str(sheet.cell(row=i, column=13).value or '').strip()
        comuna = str(sheet.cell(row=i, column=14).value or '').strip()
        cod_comuna = sheet.cell(row=i, column=15).value
        email = str(sheet.cell(row=i, column=16).value or '').strip().lower()
        telefono = sheet.cell(row=i, column=17).value
        celular = sheet.cell(row=i, column=18).value
        fecha_nac = parse_date(sheet.cell(row=i, column=19).value)
        cod_etnia = sheet.cell(row=i, column=20).value
        fecha_inc = parse_date(sheet.cell(row=i, column=21).value)
        fecha_ret = parse_date(sheet.cell(row=i, column=22).value)
        asistencia = sheet.cell(row=i, column=23).value or 0
        promedio = sheet.cell(row=i, column=24).value or 0.0

        if not run_num:
            continue

        formatted_run = format_rut(run_num, dv_run)
        student_id = f"STU-{run_num}"

        # Formatear nombre completo
        full_name = f"{ap_paterno} {ap_materno} {nombres}".strip()

        # Formatear nombre legible de grado / curso (ej: 1° Medio A)
        course_name = f"{desc_grado_raw} {letra_curso}".strip()

        # Determinar si está retirado
        is_retired = 1 if (fecha_ret is not None) else 0
        status = 'Withdrawn' if is_retired else 'Active'

        # Query de inserción o actualización
        insert_sql = """
        INSERT INTO students (
            id, run, full_name, desc_grado, status, is_retired, withdrawal_date,
            academic_year, rbd, cod_tipo_ensenanza, cod_grado, letra_curso,
            run_num, dv_run, gender, first_name, paternal_surname, maternal_surname,
            address, commune, commune_code, email, phone, mobile_phone,
            birth_date, ethnicity_code, enrollment_date, attendance_percentage, final_average
        ) VALUES (
            %s, %s, %s, %s, %s, %s, %s,
            %s, %s, %s, %s, %s,
            %s, %s, %s, %s, %s, %s,
            %s, %s, %s, %s, %s, %s,
            %s, %s, %s, %s, %s
        ) ON DUPLICATE KEY UPDATE
            run = VALUES(run),
            full_name = VALUES(full_name),
            desc_grado = VALUES(desc_grado),
            status = VALUES(status),
            is_retired = VALUES(is_retired),
            withdrawal_date = VALUES(withdrawal_date),
            email = VALUES(email),
            mobile_phone = VALUES(mobile_phone),
            address = VALUES(address)
        """

        vals = (
            student_id, formatted_run, full_name, course_name, status, is_retired, fecha_ret,
            year or 2026, rbd or 3941, cod_tipo_ens, cod_grado, letra_curso,
            str(run_num), str(dv_run or ''), gender, nombres, ap_paterno, ap_materno,
            direccion if direccion else None, comuna if comuna else None, str(cod_comuna or ''),
            email if email else None, str(telefono or '') if telefono else None, str(celular or '') if celular else None,
            fecha_nac, cod_etnia, fecha_inc, asistencia, promedio
        )

        cursor.execute(insert_sql, vals)
        inserted_count += 1

    db.commit()
    cursor.close()
    db.close()

    print(f"🎉 ¡Importación completada! Se procesaron {inserted_count} estudiantes exitosamente en XAMPP MySQL.")

if __name__ == '__main__':
    import_roster()
